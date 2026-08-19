"""Genereert dataverzoek-e-mails aan leveranciers.

- Verzamelt de ontbrekende compliance-velden per leverancier (gegroepeerd per
  wetgeving).
- Bouwt een Excel-bijlage met die ontbrekende velden.
- Genereert de mailtekst via de Anthropic API (model: claude-sonnet-4-6) op
  basis van de ontbrekende data en wetgeving. Zonder API-sleutel (of bij een
  fout) valt het terug op een net sjabloon, zodat de functionaliteit altijd
  werkt.
"""
import io
import logging
import os
from collections import defaultdict
from datetime import date
from typing import Optional

from sqlalchemy.orm import Session

from . import models, compliance, veld_vertaling

log = logging.getLogger(__name__)

CC_ADRES = "compliance@uwbedrijf.nl"
PORTAAL_BASIS = "https://portaal.powercompliance.nl/leverancier"
MODEL = "claude-sonnet-4-6"

# GPSR (EU 2023/988) vereist een verantwoordelijke marktdeelnemer in de EU met
# naam én adres. Ontbreken die bedrijfsgegevens bij de leverancier zelf, dan
# vragen we ze mee uit onder de GPSR-noemer. (attribuut, NL-label, EN-label).
GPSR_CODE = "GPSR"
LEVERANCIER_CONTACTVELDEN = [
    ("naam", "Bedrijfsnaam", "Company name"),
    ("adres", "Postadres (straat + huisnummer)", "Postal address (street + number)"),
    ("postcode", "Postcode", "Postal code"),
    ("stad", "Plaats", "City"),
    ("land", "Land", "Country"),
    ("telefoon", "Telefoonnummer", "Phone number"),
    ("email", "E-mailadres", "Email address"),
    ("kvk_nummer", "KvK-nummer", "Chamber of Commerce number"),
    ("btw_nummer", "BTW-nummer", "VAT number"),
]


def ontbrekende_leverancier_gegevens(
    leverancier: models.Leverancier, taal: str = "nl"
) -> list:
    """Geef de labels van de bedrijfscontactgegevens die bij de leverancier nog
    ontbreken (leeg of niet ingevuld). Gekoppeld aan de GPSR-vereiste voor de
    verantwoordelijke EU-marktdeelnemer."""
    ontbrekend = []
    for attr, nl_label, en_label in LEVERANCIER_CONTACTVELDEN:
        waarde = getattr(leverancier, attr, None)
        if waarde is None or (isinstance(waarde, str) and not waarde.strip()):
            ontbrekend.append(en_label if taal == "en" else nl_label)
    return ontbrekend


# ---------- data verzamelen ----------
def verzamel_ontbrekend(
    db: Session,
    leverancier: models.Leverancier,
    wetgeving_code: Optional[str] = None,
    product_id: Optional[int] = None,
    taal: str = "nl",
):
    """Geef (per_product, per_wetgeving) terug.

    per_product: lijst van (product, [ComplianceVeld]) met ontbrekende velden.
    per_wetgeving: dict wetgeving_code -> set van veldnamen (Engels bij taal="en").

    Scoping:
    - ``product_id``: beperk tot dat ene product (uitvraag voor 1 product).
    - ``wetgeving_code``: alleen ontbrekende velden van die ene wetgeving.
    - geen van beide: alle ontbrekende velden over alle producten van de leverancier.
    """
    per_product = []
    per_wet = defaultdict(set)
    # Bij product-scope alleen dat product bekijken (mits het van deze leverancier is).
    producten = leverancier.producten
    if product_id is not None:
        producten = [p for p in producten if p.id == product_id]
    for product in producten:
        # Ontbrekende velden = velden zonder (niet-lege) waarde voor dit product,
        # bepaald via de LEFT JOIN-query in compliance.ontbrekende_velden_voor_product.
        ontbrekend = compliance.ontbrekende_velden_voor_product(db, product)
        if wetgeving_code:
            ontbrekend = [
                v
                for v in ontbrekend
                if v.wetgeving and v.wetgeving.code == wetgeving_code
            ]
        if ontbrekend:
            per_product.append((product, ontbrekend))
            for v in ontbrekend:
                code = v.wetgeving.code if v.wetgeving else "—"
                per_wet[code].add(veld_vertaling.veld_naam(v, taal))
    # Leverancier-brede bedrijfsgegevens (GPSR: verantwoordelijke EU-marktdeelnemer).
    # Alleen op leverancier-scope meenemen (niet bij een uitvraag voor één product),
    # en alleen als de uitvraag GPSR omvat (geen filter, of expliciet GPSR).
    if product_id is None and wetgeving_code in (None, GPSR_CODE):
        prefix = "Company details: " if taal == "en" else "Bedrijfsgegevens: "
        for label in ontbrekende_leverancier_gegevens(leverancier, taal):
            per_wet[GPSR_CODE].add(prefix + label)
    totaal = sum(len(v) for _, v in per_product)
    log.info(
        "verzamel_ontbrekend(leverancier=%r, wetgeving=%s, product_id=%s): %d producten met samen %d ontbrekende velden",
        leverancier.naam, wetgeving_code, product_id, len(per_product), totaal,
    )
    return per_product, dict(per_wet)


def leveranciers_met_ontbrekend_voor_wetgeving(db: Session, wetgeving_code: str):
    """Alle leveranciers die ontbrekende data hebben voor deze wetgeving."""
    result = []
    leveranciers = db.query(models.Leverancier).order_by(models.Leverancier.naam).all()
    for lev in leveranciers:
        per_product, _ = verzamel_ontbrekend(db, lev, wetgeving_code)
        aantal_velden = sum(len(velden) for _, velden in per_product)
        # Bij GPSR tellen ook ontbrekende bedrijfsgegevens mee, zodat een
        # leverancier met alleen incomplete contactdata óók wordt uitgevraagd.
        if wetgeving_code == GPSR_CODE:
            aantal_velden += len(ontbrekende_leverancier_gegevens(lev))
        if aantal_velden > 0:
            result.append(
                {
                    "id": lev.id,
                    "naam": lev.naam,
                    "contactpersoon": lev.contactpersoon,
                    "email": lev.email,
                    "aantal_velden": aantal_velden,
                    "aantal_producten": len(per_product),
                }
            )
    return result


def portaal_link(leverancier: models.Leverancier) -> str:
    return f"{PORTAAL_BASIS}/{leverancier.id}/aanleveren"


def _veilige_naam(tekst: str, fallback: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in tekst).strip("_") or fallback


def bijlage_naam(
    leverancier: models.Leverancier,
    wetgeving_code: Optional[str] = None,
    product: Optional[models.Product] = None,
) -> str:
    veilig = _veilige_naam(leverancier.naam, str(leverancier.id))
    delen = ["ontbrekende_data", veilig]
    if product is not None:
        delen.append(_veilige_naam(product.artikelnummer or product.naam, f"p{product.id}"))
    if wetgeving_code:
        delen.append(wetgeving_code)
    return "_".join(delen) + ".xlsx"


# ---------- Excel-bijlage ----------
def bouw_excel(
    db: Session,
    leverancier: models.Leverancier,
    wetgeving_code: Optional[str] = None,
    product_id: Optional[int] = None,
    taal: str = "nl",
) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    per_product, _ = verzamel_ontbrekend(
        db, leverancier, wetgeving_code, product_id, taal
    )
    wb = Workbook()
    ws = wb.active
    ws.title = veld_vertaling.EXCEL_TITEL.get(taal, veld_vertaling.EXCEL_TITEL["nl"])
    kop = veld_vertaling.EXCEL_KOPPEN.get(taal, veld_vertaling.EXCEL_KOPPEN["nl"])
    ws.append(kop)
    for cel in ws[1]:
        cel.font = Font(bold=True)
    for product, velden in per_product:
        for v in velden:
            ws.append(
                [
                    product.naam,
                    product.artikelnummer or "",
                    product.ean or "",
                    v.wetgeving.code if v.wetgeving else "",
                    veld_vertaling.veld_naam(v, taal),
                    veld_vertaling.veld_type(v.veld_type, taal),
                    "",
                ]
            )
    breedtes = [28, 18, 16, 12, 32, 12, 24]
    for i, b in enumerate(breedtes, start=1):
        ws.column_dimensions[chr(64 + i)].width = b

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    aantal_regels = sum(len(velden) for _, velden in per_product)
    log.info(
        "bouw_excel(leverancier=%r, wetgeving=%s, product_id=%s): %d datarijen, %d bytes",
        leverancier.naam, wetgeving_code, product_id, aantal_regels, buf.getbuffer().nbytes,
    )
    return buf


def bouw_excel_bytes(
    db: Session,
    leverancier: models.Leverancier,
    wetgeving_code: Optional[str] = None,
    product_id: Optional[int] = None,
    taal: str = "nl",
) -> bytes:
    """Zoals bouw_excel, maar geeft de rauwe bytes terug — handig om als
    e-mailbijlage mee te sturen."""
    return bouw_excel(db, leverancier, wetgeving_code, product_id, taal).getvalue()


# ---------- onderwerp ----------
def maak_onderwerp(
    leverancier: models.Leverancier,
    per_wet: dict,
    taal: str,
    deadline: Optional[date] = None,
    product: Optional[models.Product] = None,
) -> str:
    """Onderwerp met leverancier (+ product) + wetgeving + (optioneel) deadline."""
    codes = ", ".join(sorted(per_wet.keys())) if per_wet else ""
    onderwerp_van = leverancier.naam
    if product is not None:
        onderwerp_van = f"{leverancier.naam} – {product.naam}"
    if taal == "en":
        s = f"{onderwerp_van} – missing product compliance data"
    else:
        s = f"{onderwerp_van} – ontbrekende productcompliance-data"
    if codes:
        s += f" ({codes})"
    if deadline:
        s += f" – deadline {deadline.isoformat()}"
    return s


# ---------- mailtekst (AI + fallback) ----------
def _samenvatting_tekst(per_wet: dict) -> str:
    regels = []
    for code, velden in sorted(per_wet.items()):
        regels.append(f"- {code}: {', '.join(sorted(velden))}")
    return "\n".join(regels)


SYSTEM_PROMPT = (
    "Je bent een compliance-medewerker bij een groothandel. Je schrijft korte, "
    "professionele en vriendelijke zakelijke e-mails aan leveranciers met het "
    "verzoek ontbrekende productdata aan te leveren in het kader van EU-wetgeving. "
    "Geef ALLEEN de tekst van de e-mail (begroeting t/m afsluiting), zonder "
    "onderwerpregel en zonder uitleg eromheen."
)


def _bouw_prompt(
    leverancier: models.Leverancier,
    per_wet: dict,
    deadline: Optional[date],
    taal: str,
    link: str,
    aantal_velden: int,
    aantal_producten: int,
) -> str:
    samenvatting = _samenvatting_tekst(per_wet)
    contact = leverancier.contactpersoon or ("Sir/Madam" if taal == "en" else "heer/mevrouw")
    deadline_str = deadline.isoformat() if deadline else None

    if taal == "en":
        return f"""Write a professional business email in ENGLISH to a supplier.

Supplier: {leverancier.naam}
Contact person: {contact}
Number of products with missing data: {aantal_producten}
Total missing fields: {aantal_velden}

Missing data, grouped by EU legislation:
{samenvatting}

Requirements for the email:
- Briefly explain that we need this data to stay compliant with the listed EU regulations (e.g. PPWR, Battery Regulation, REACH/CLP, CPR, GPSR, ErP).
- Mention that an Excel file listing the exact missing fields is attached.
- Offer the supplier TWO ways to deliver the data:
  1. Reply to this email with the data as plain text.
  2. Upload the data via our portal: {link}
- {"Ask them to deliver before the deadline: " + deadline_str if deadline_str else "Politely ask them to respond as soon as possible."}
- Keep it concise and friendly. Sign off as "The Compliance Team".
Return only the email body text."""

    return f"""Schrijf een professionele zakelijke e-mail in het NEDERLANDS aan een leverancier.

Leverancier: {leverancier.naam}
Contactpersoon: {contact}
Aantal producten met ontbrekende data: {aantal_producten}
Totaal ontbrekende velden: {aantal_velden}

Ontbrekende data, gegroepeerd per EU-wetgeving:
{samenvatting}

Eisen aan de e-mail:
- Leg kort uit dat we deze data nodig hebben om te voldoen aan de genoemde EU-wetgeving (zoals PPWR, Batterijverordening, REACH/CLP, CPR, GPSR, ErP).
- Vermeld dat een Excel-bestand met de exacte ontbrekende velden is bijgevoegd.
- Bied de leverancier TWEE manieren om de data aan te leveren:
  1. Reageer op deze e-mail met de data als platte tekst.
  2. Upload de data via ons portaal: {link}
- {"Vraag om aanlevering vóór de deadline: " + deadline_str if deadline_str else "Vraag vriendelijk om zo snel mogelijk te reageren."}
- Houd het beknopt en vriendelijk. Onderteken met "Het Compliance-team".
Geef alleen de tekst van de e-mail terug."""


def _fallback_tekst(
    leverancier: models.Leverancier,
    per_wet: dict,
    deadline: Optional[date],
    taal: str,
    link: str,
) -> str:
    samenvatting = _samenvatting_tekst(per_wet)
    contact = leverancier.contactpersoon
    deadline_str = deadline.isoformat() if deadline else None
    if taal == "en":
        aanhef = f"Dear {contact}," if contact else "Dear Sir/Madam,"
        deadline_zin = (
            f"Please deliver the data before {deadline_str}."
            if deadline_str
            else "Please respond at your earliest convenience."
        )
        return f"""{aanhef}

To keep our shared product range compliant with the relevant EU regulations, we are missing some product data from you. The missing fields, grouped by legislation, are:

{samenvatting}

A spreadsheet listing the exact missing fields is attached to this email.

You can deliver the data in either of two ways:
1. Reply to this email with the data as plain text.
2. Upload the data via our portal: {link}

{deadline_zin}

Thank you in advance for your cooperation.

Kind regards,
The Compliance Team"""
    aanhef = f"Beste {contact}," if contact else "Geachte heer/mevrouw,"
    deadline_zin = (
        f"Graag ontvangen wij de data vóór {deadline_str}."
        if deadline_str
        else "Wij ontvangen de data graag zo spoedig mogelijk."
    )
    return f"""{aanhef}

Om ons gezamenlijke productassortiment te laten voldoen aan de relevante EU-wetgeving, ontbreekt bij ons nog een aantal productgegevens van uw kant. De ontbrekende velden, gegroepeerd per wetgeving, zijn:

{samenvatting}

In de bijlage van deze e-mail vindt u een overzicht (Excel) met de exacte ontbrekende velden.

U kunt de data op twee manieren aanleveren:
1. Reageer op deze e-mail met de gegevens als platte tekst.
2. Upload de gegevens via ons portaal: {link}

{deadline_zin}

Alvast hartelijk dank voor uw medewerking.

Met vriendelijke groet,
Het Compliance-team"""


# ---------- Placeholders & sjablonen (voor sequence-stap mailinhoud) ----------
# Placeholders die de beheerder in een eigen stap-mailtekst/onderwerp mag
# gebruiken; ze worden bij verzending én in de preview per leverancier ingevuld.
PLACEHOLDERS = [
    "{aanhef}",
    "{leverancier}",
    "{contactpersoon}",
    "{ontbrekende_data}",
    "{portaal_link}",
    "{aantal_velden}",
    "{aantal_producten}",
]


def placeholder_waarden(
    leverancier: Optional[models.Leverancier],
    per_wet: dict,
    link: str,
    aantal_velden: int,
    aantal_producten: int,
    taal: str = "nl",
) -> dict:
    """Concrete waarden voor de placeholders van een sequence-stap."""
    contact = leverancier.contactpersoon if leverancier else None
    naam = leverancier.naam if leverancier else "Voorbeeld Leverancier B.V."
    if taal == "en":
        aanhef = f"Dear {contact}," if contact else "Dear Sir/Madam,"
        onbekend = "Sir/Madam"
    else:
        aanhef = f"Beste {contact}," if contact else "Geachte heer/mevrouw,"
        onbekend = "heer/mevrouw"
    return {
        "aanhef": aanhef,
        "leverancier": naam,
        "contactpersoon": contact or onbekend,
        "ontbrekende_data": _samenvatting_tekst(per_wet),
        "portaal_link": link,
        "aantal_velden": str(aantal_velden),
        "aantal_producten": str(aantal_producten),
    }


def render_sjabloon(tekst: str, waarden: dict) -> str:
    """Vervang {placeholder}-tokens in tekst. Onbekende tokens blijven staan."""
    if not tekst:
        return tekst
    for sleutel, waarde in waarden.items():
        tekst = tekst.replace("{" + sleutel + "}", waarde or "")
    return tekst


def voorbeeld_context(db: Session, wetgeving_code: Optional[str] = None):
    """Zoek een representatieve leverancier met ontbrekende data voor preview/
    generatie. Geeft (leverancier|None, per_wet, aantal_velden, aantal_producten).

    Zonder kandidaat wordt een fictief voorbeeld teruggegeven zodat de preview
    altijd iets toont."""
    leveranciers = db.query(models.Leverancier).order_by(models.Leverancier.naam).all()
    for lev in leveranciers:
        per_product, per_wet = verzamel_ontbrekend(db, lev, wetgeving_code)
        if per_product:
            aantal_velden = sum(len(v) for _, v in per_product)
            return lev, per_wet, aantal_velden, len(per_product)
    # fictief voorbeeld
    code = wetgeving_code or "PPWR"
    per_wet = {code: {"Voorbeeldveld A", "Voorbeeldveld B"}}
    return None, per_wet, 2, 1


def _sjabloon_fallback(taal: str) -> str:
    """Herbruikbaar mailsjabloon met placeholders (fallback zonder AI)."""
    if taal == "en":
        return (
            "{aanhef}\n\n"
            "To keep our shared product range compliant with the relevant EU "
            "regulations, we are still missing some product data from you. The "
            "missing fields, grouped by legislation, are:\n\n"
            "{ontbrekende_data}\n\n"
            "A spreadsheet listing the exact missing fields is attached to this email.\n\n"
            "You can deliver the data in either of two ways:\n"
            "1. Reply to this email with the data as plain text.\n"
            "2. Upload the data via our portal: {portaal_link}\n\n"
            "Thank you in advance for your cooperation.\n\n"
            "Kind regards,\nThe Compliance Team"
        )
    return (
        "{aanhef}\n\n"
        "Om ons gezamenlijke productassortiment te laten voldoen aan de relevante "
        "EU-wetgeving, ontbreekt bij ons nog een aantal productgegevens van uw kant. "
        "De ontbrekende velden, gegroepeerd per wetgeving, zijn:\n\n"
        "{ontbrekende_data}\n\n"
        "In de bijlage van deze e-mail vindt u een overzicht (Excel) met de exacte "
        "ontbrekende velden.\n\n"
        "U kunt de data op twee manieren aanleveren:\n"
        "1. Reageer op deze e-mail met de gegevens als platte tekst.\n"
        "2. Upload de gegevens via ons portaal: {portaal_link}\n\n"
        "Alvast hartelijk dank voor uw medewerking.\n\n"
        "Met vriendelijke groet,\nHet Compliance-team"
    )


def _sjabloon_onderwerp(wetgeving_code: Optional[str], taal: str) -> str:
    scope = f" ({wetgeving_code})" if wetgeving_code else ""
    if taal == "en":
        return "{leverancier} – missing product compliance data" + scope
    return "{leverancier} – ontbrekende productcompliance-data" + scope


def genereer_sjabloon(wetgeving_code: Optional[str], taal: str):
    """Genereer een herbruikbaar mailSJABLOON met placeholders via de Anthropic
    API. Geeft (onderwerp, tekst, ai_gebruikt, ai_fout) terug. Valt terug op een
    net sjabloon zonder API-sleutel of bij een fout."""
    onderwerp = _sjabloon_onderwerp(wetgeving_code, taal)
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return (
            onderwerp,
            _sjabloon_fallback(taal),
            False,
            "ANTHROPIC_API_KEY niet ingesteld — sjabloontekst gebruikt.",
        )
    scope_zin = (
        f" over de wetgeving {wetgeving_code}" if wetgeving_code else " over EU-compliancedata"
    )
    if taal == "en":
        scope_zin = f" about {wetgeving_code}" if wetgeving_code else " about EU compliance data"
        prompt = f"""Write a REUSABLE email TEMPLATE in ENGLISH for a recurring data request to suppliers{scope_zin}.
Use EXACTLY these placeholder tokens verbatim (do not translate, keep the curly braces) at the appropriate spots:
- {{aanhef}}  (will be replaced by e.g. "Dear Jan,")
- {{ontbrekende_data}}  (a list of missing fields grouped per regulation)
- {{portaal_link}}  (a link to the supplier portal)

Requirements: briefly explain why the data is needed, mention that an Excel attachment lists the exact missing fields, offer two delivery options (reply as plain text, or upload via {{portaal_link}}), keep it concise and friendly, and sign off as "The Compliance Team". Return only the template body text."""
    else:
        prompt = f"""Schrijf een HERBRUIKBAAR e-mailSJABLOON in het NEDERLANDS voor een terugkerend dataverzoek aan leveranciers{scope_zin}.
Gebruik EXACT deze placeholder-tokens letterlijk (niet vertalen, accolades behouden) op de juiste plek:
- {{aanhef}}  (wordt vervangen door bv. "Beste Jan,")
- {{ontbrekende_data}}  (een lijst met ontbrekende velden gegroepeerd per wetgeving)
- {{portaal_link}}  (een link naar het leveranciersportaal)

Eisen: leg kort uit waarom de data nodig is, vermeld dat een Excel-bijlage de exacte ontbrekende velden bevat, bied twee aanleverwijzen (reageren als platte tekst, of uploaden via {{portaal_link}}), houd het beknopt en vriendelijk, en onderteken met "Het Compliance-team". Geef alleen de sjabloontekst terug."""
    try:
        import anthropic

        client = anthropic.Anthropic(api_key=api_key)
        resp = client.messages.create(
            model=MODEL,
            max_tokens=1500,
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": prompt}],
        )
        tekst = "".join(b.text for b in resp.content if b.type == "text").strip()
        if not tekst:
            raise ValueError("Leeg antwoord van de API")
        return onderwerp, tekst, True, None
    except Exception as e:  # noqa: BLE001
        return (
            onderwerp,
            _sjabloon_fallback(taal),
            False,
            f"AI-generatie mislukt ({type(e).__name__}): {e}",
        )


def genereer_tekst(
    leverancier: models.Leverancier,
    per_wet: dict,
    deadline: Optional[date],
    taal: str,
    link: str,
    aantal_velden: int,
    aantal_producten: int,
):
    """Geef (tekst, ai_gebruikt, ai_fout) terug."""
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return (
            _fallback_tekst(leverancier, per_wet, deadline, taal, link),
            False,
            "ANTHROPIC_API_KEY niet ingesteld — sjabloontekst gebruikt.",
        )
    try:
        import anthropic

        client = anthropic.Anthropic(api_key=api_key)
        prompt = _bouw_prompt(
            leverancier, per_wet, deadline, taal, link, aantal_velden, aantal_producten
        )
        resp = client.messages.create(
            model=MODEL,
            max_tokens=1500,
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": prompt}],
        )
        tekst = "".join(b.text for b in resp.content if b.type == "text").strip()
        if not tekst:
            raise ValueError("Leeg antwoord van de API")
        return tekst, True, None
    except Exception as e:  # noqa: BLE001 — terugvallen op sjabloon bij elke fout
        return (
            _fallback_tekst(leverancier, per_wet, deadline, taal, link),
            False,
            f"AI-generatie mislukt ({type(e).__name__}): {e}",
        )
